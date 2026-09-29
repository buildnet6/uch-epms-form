"""Comments, quarterly scoring and dates for an EPMS workbook, written from the nurse's own figures.

Everything here is derived from what the nurse submitted and the contract's own rating scale
(O 100%, E 90%, VG 80%, G 70%, F 60%, P 50% of a task's marks). The appraiser and counter-signing officer
still read, amend and sign; these are sensible defaults so no box is left empty.
"""
import re
from datetime import datetime
from collections import Counter
from openpyxl.styles import Alignment

BAND_NAMES = ['Outstanding', 'Excellent', 'Very good', 'Good', 'Fair', 'Poor']
BAND_MARKS = [100, 90, 80, 70, 60, 50]
QUARTER_END = [(3, 31), (6, 30), (9, 30), (12, 31)]
ISSUE_WORDS = [('shortage of personnel', 'shortage of personnel'), ('shortage of manpower', 'shortage of personnel'),
               ('poor power supply', 'poor power supply'), ('irregular power supply', 'poor power supply'),
               ('power outage', 'poor power supply'), ('high patient load', 'the high patient load'),
               ('shortage of consumables', 'shortage of consumables'), ('equipment breakdown', 'equipment breakdown'),
               ('network downtime', 'network downtime'), ('transport', 'limited transport')]
WRAP = Alignment(wrap_text=True, vertical='center', horizontal='left')


def _task_cls():
    from summary_pdf import Task, short_task      # imported late: summary_pdf imports fill_epms
    return Task, short_task


# plain names for tasks in comments ("met target in hand hygiene compliance and clinic workflow")
PLAIN = [('formulate goals for and direct', 'sectional goals and operations'), ('participates in direct holistic patient care', 'holistic patient care'),
         ('assess, plans, develops', 'planning and monitoring of patient care'),
         ('compliance with hand hygiene', 'hand hygiene compliance'), ('prompt response to patients', "prompt response to patients' needs"),
         ('community eye screening', 'community eye screening'), ('clinic workflow', 'clinic workflow'),
         ('strengthen administrative coordination', 'administrative coordination'), ('strengthen clinical documentation', 'clinical documentation'),
         ('develop and publish a triage sop', 'the triage SOP'), ('disinfection of equipment', 'disinfection and sterilization'),
         ('give reliable health information', 'health information to clients'), ('remind specialists', 'appointment reminders'),
         ('giving health information to mothers', 'health information to mothers'), ('safe delivery of babies', 'safe delivery and infection prevention'),
         ('cascade of pms', 'cascading the PMS process'), ('compliance with ecm', 'ECM compliance'),
         ('through the usage of the staggered appointment', 'the staggered appointment system'), ('provide, continuous training', 'staff training'),
         ('coordination of the promotion', 'the promotion exercise'), ('capacity building', 'capacity building'),
         ('task shifting', 'task shifting to allied staff'), ('optimize operating theatre', 'theatre scheduling'),
         ('introduce preventive maintenance', 'preventive maintenance'), ('daily health talk', 'daily health talks'),
         ('document attendance', 'attendance documentation'), ('regular training of staff', 'staff training')]


def short(text, n=48):
    low = ' '.join(str(text or '').split()).lower()
    for key, plain in PLAIN:
        if low.startswith(key):
            return plain
    _, short_task = _task_cls()
    t = short_task(text, n).rstrip('…').strip()
    return t[0].lower() + t[1:] if t[:2].isupper() is False and t[:1].isupper() and not t[1:2].isupper() else t


def cap(t):
    return t[0].upper() + t[1:] if t else t


def join(items):
    items = [i for i in dict.fromkeys(items) if i]
    if len(items) <= 1:
        return ''.join(items)
    return ', '.join(items[:-1]) + ' and ' + items[-1]


def fmt(v):
    if isinstance(v, float):
        v = round(v, 2)
        return str(int(v)) if v.is_integer() else f"{v:g}"
    return str(v)


def main_issue(texts):
    c = Counter()
    for t in texts:
        low = (t or '').lower()
        for key, label in ISSUE_WORDS:
            if key in low:
                c[label] += 1
    return c.most_common(1)[0][0] if c else ''


class Entry:
    def __init__(self, row, task, value):
        self.row, self.task, self.value = row, task, value
        self.name = short(row['kra'])
        self.rating = task.rating(value) if task and isinstance(value, (int, float)) else None
        self.meets = task.meets(value) if task and isinstance(value, (int, float)) else None
        if self.rating is None and self.meets is not None:          # rows with no usable scale: judge against target
            self.rating = 'Excellent' if self.meets else 'Good'

    @property
    def good(self):
        return self.meets or self.rating in ('Outstanding', 'Excellent')

    def against(self):
        t = self.task
        if not t or not isinstance(t.target, (int, float)):
            return fmt(self.value)
        if t.frac:
            return f"{fmt(round(self.value * 100, 1))}% against a target of {fmt(t.target * 100)}%"
        unit = '%' if t.unit == '%' else (' minutes' if t.unit.lower().startswith('min') else '')
        return f"{fmt(self.value)}{unit} against a target of {fmt(t.target)}{unit}"


def pick(options, seed):
    return options[seed % len(options)]


# ---------------- monthly ----------------
def monthly_comments(entries, issues, mi, first_name):
    good = [e for e in entries if e.good]
    weak = [e for e in entries if e.meets is False and not e.good]
    issue = main_issue(issues)
    if not entries:
        return None, None
    # appraiser (supervisor)
    if not weak:
        appraiser = pick([
            f"Excellent performance this month. Targets on {join([e.name for e in good])} were met. Keep it up.",
            f"Very commendable. {cap(join([e.name for e in good]))} achieved as planned. Maintain this standard.",
            f"Good work this month; all set targets were met. Keep sustaining the standard of care.",
        ], mi)
    elif good:
        w = weak[0]
        appraiser = pick([
            f"Good effort. {cap(join([e.name for e in good]))} met target; {w.name} ({w.against()}) needs more attention next month.",
            f"Satisfactory performance. Well done on {join([e.name for e in good])}. Work on improving {w.name} ({w.against()}).",
        ], mi)
    else:
        w = weak[0]
        appraiser = pick([
            f"Fair performance. {cap(w.name)} ({w.against()}) fell below target. Let us work together to improve next month.",
            f"Below target this month on {join([e.name for e in weak])}. More effort is needed; support will be provided.",
        ], mi)
    # appraisee
    if issue:
        appraisee = pick([
            f"I agree with this review. {issue[0].upper() + issue[1:]} affected my work this month, but I remain committed to meeting my targets.",
            f"I agree. Despite {issue}, I did my best and will keep improving.",
            f"I accept this review. I will continue to give my best in spite of {issue}.",
        ], mi)
    else:
        appraisee = pick(["I agree with this review and will sustain my performance.",
                          "I agree with this review. I will keep up the good work.",
                          "I accept this review and remain committed to my targets."], mi)
    return appraisee, appraiser


# ---------------- quarterly ----------------
def level_phrase(pct):
    if pct >= 90: return 'Excellent'
    if pct >= 80: return 'Very good'
    if pct >= 70: return 'Good'
    if pct >= 60: return 'Fair'
    return 'Poor'


def fill_quarter(ws, qi, ks, entries_by_key, contract, pms, data, year):
    """ks: task rows already written at 19.. by fill(); entries_by_key: {(code, norm kra): [Entry,...]} for the quarter."""
    from fill_epms import norm
    Task, _ = _task_cls()
    extras = data.get('extras') or {}
    q_entries = []
    for i, row in enumerate(ks):
        r = 19 + i
        es = entries_by_key.get((row['code'], norm(row['kra'])), [])
        vals = [e.value for e in es if isinstance(e.value, (int, float))]
        cr = contract.find(row['code'], row['kra'])
        task = Task(pms, cr) if cr else None
        if not vals:
            continue
        ach = sum(vals) / len(vals)
        ach = round(ach, 2) if task and task.frac else round(ach, 1)
        e = Entry(row, task, ach)
        q_entries.append(e)
        ws.cell(r, 16).value = int(ach) if float(ach).is_integer() else ach                 # P achievement
        if task and task.frac and ach <= 1:
            ws.cell(r, 16).number_format = '0%'
        marks = BAND_MARKS[BAND_NAMES.index(e.rating)] if e.rating in BAND_NAMES else None
        ws.cell(r, 17).value = marks                                                          # Q raw score
        ws.cell(r, 18).value = f'=IF(Q{r}="","",F{r}*Q{r}/100)' if marks is not None else None   # R weighted raw
        for c in (16, 17, 18):
            ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')
    # performance level for the quarter (share of available marks)
    marks_list = [BAND_MARKS[BAND_NAMES.index(e.rating)] for e in q_entries if e.rating in BAND_NAMES]
    level = (sum(marks_list) / len(marks_list)) if marks_list else 80

    # Section 5 competencies (20 marks) and Section 6 operations (10 marks)
    comp_rows = {30: 104, 31: 105, 32: 106, 34: 108, 35: 109, 36: 110, 38: 112, 39: 113, 40: 114}
    targets = {30: 4, 31: 3, 32: 3, 34: 2, 35: 2, 36: 1, 38: 2, 39: 2, 40: 1}
    deduct = 0 if level >= 95 else 1 if level >= 88 else 2 if level >= 80 else 3
    order = [31, 36, 32][:deduct]
    for qr, cr in comp_rows.items():
        ws.cell(qr, 3).value = ' '.join(str(pms.cell(cr, 6).value or '').split()) or None
        ws.cell(qr, 3).alignment = WRAP
        ws.cell(qr, 8).value = targets[qr]
        ws.cell(qr, 9).value = targets[qr] - (1 if qr in order and targets[qr] > 1 else 0)
    for qr, cr in {42: 116, 43: 117, 44: 118}.items():
        ws.cell(qr, 3).value = ' '.join(str(pms.cell(cr, 6).value or '').split()) or None
        ws.cell(qr, 3).alignment = WRAP
    ws['I44'] = 3 if level >= 85 else 2
    for qr in list(comp_rows) + [42, 43, 44]:
        for c in (8, 9):
            ws.cell(qr, c).alignment = Alignment(horizontal='center', vertical='center')

    # Section 6 summary
    ws['O47'] = '=R25'; ws['O48'] = '=I41'; ws['O49'] = '=I45'
    for a in ('O47', 'O48', 'O49', 'O50'):
        ws[a].number_format = '0.0'; ws[a].alignment = Alignment(horizontal='center', vertical='center')

    # strengths / areas for improvement / comments
    good = [e for e in q_entries if e.good]
    weak = [e for e in q_entries if not e.good]
    issues = [e.row.get('issues', '') for es in entries_by_key.values() for e in es]
    issue = main_issue(issues)
    phrase = level_phrase(level)
    strengths = (f"Met or exceeded target in {join([e.name for e in good])}. " if good else "") + pick([
        "Dedicated, reliable and works well with the team.",
        "Committed to patient care and dependable under pressure.",
        "Punctual, cooperative and shows good leadership on the unit.",
        "Consistent, hardworking and supportive of colleagues."], qi)
    if weak:
        improve = f"Improve on {join([f'{e.name} (averaged {e.against()})' for e in weak])}."
    else:
        improve = "Sustain the current standard across all tasks."
    tn = re.sub(r'\s+(is\s+)?(required|needed)$', '', (extras.get('training_needs') or '').strip().rstrip('.'), flags=re.I)
    if tn:
        improve += f" Would benefit from {tn[0].lower() + tn[1:]}."
    appraisee = pick([
        "I agree with this appraisal and will keep working to improve" + (f", especially on {weak[0].name}." if weak else "."),
        "I accept this appraisal. Thank you for the support; I will sustain and improve my performance.",
        "I agree with the assessment" + (f" and will address {issue} as best I can." if issue else " and remain committed to my targets."),
        "I agree. I will continue to give my best in the coming quarter."], qi)
    supervisor = f"{phrase} performance this quarter. " + (
        pick(["Keep up the good work.", "The officer is commended.", "Maintain this standard.", "Well done."], qi) if not weak
        else f"More attention is needed on {weak[0].name}.")
    cso = pick([
        "I agree with the supervisor's assessment.",
        "I concur with the appraiser's assessment. The officer is encouraged to keep improving.",
        "Assessment noted and endorsed.",
        "I agree with the supervisor. The officer should sustain this performance."], qi) + (
        "" if not weak or qi % 2 else " Support should be given to improve on the areas identified.")

    ws['B52'] = strengths; ws['B54'] = improve
    own = (((data.get('quarterly') or [{}] * 4)[qi] or {}).get('appraisee_comment') or '').strip() if len(data.get('quarterly') or []) > qi else ''
    ws['C55'] = own or appraisee
    ws['E55'] = "Supervisor's Comment\n" + supervisor
    ws['G55'] = "Counter Supervisor's Comment"
    ws['H55'] = cso
    for a in ('B52', 'B54', 'C55', 'H55'):
        ws[a].alignment = WRAP
    ws['E55'].alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    # dates: the quarter's last day
    m, d = QUARTER_END[qi]
    for a in ('C62', 'F62', 'H62'):
        ws[a] = datetime(year, m, d)
        ws[a].number_format = 'dd/mm/yyyy'
        ws[a].alignment = Alignment(horizontal='center', vertical='center')
    return q_entries


# ---------------- contract ----------------
def contract_comments(pms):
    pms['C120'] = "I agree with the tasks and targets set in this contract."
    pms['I120'] = "The tasks and targets were discussed and agreed with the officer. They are realistic and achievable."
    pms['P120'] = "Reviewed and approved."
    for a in ('C120', 'I120', 'P120'):
        pms[a].alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')


def monthly_entries(month_rows, contract, pms):
    Task, _ = _task_cls()
    out = []
    for mrows in month_rows:
        es = []
        for row in mrows:
            if not row['kra']:
                continue
            from fill_epms import num_val
            v = num_val(row['output'])
            cr = contract.find(row['code'], row['kra'])
            es.append(Entry(row, Task(pms, cr) if cr else None, v if isinstance(v, (int, float)) else None))
        out.append(es)
    return out
