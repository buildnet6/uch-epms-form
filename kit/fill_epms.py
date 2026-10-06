"""Fill a nurse's EPMS workbook from one submission of the online EPMS details form.

Usage:  python fill_epms.py submission.json [more.json ...] --out OUTDIR
Each JSON file is the `data` object saved by the form (form == "UCH-EPMS-DETAILS").
"""
import sys, os, re, json, calendar, argparse
from copy import copy
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import column_index_from_string, get_column_letter


def place_signature_fit(ws, sig_path, rng, max_h=54, pad=5):
    """Place a signature inside the merged box `rng` (e.g. 'I124:M126'), as large as fits up to max_h pixels."""
    from openpyxl.utils.cell import range_boundaries
    c1, r1, c2, r2 = range_boundaries(rng)
    dw = ws.sheet_format.defaultColWidth or 8.43; dh = ws.sheet_format.defaultRowHeight or 15
    def cw(c):                                   # read widths without creating column entries (they change wrapping)
        d = ws.column_dimensions.get(get_column_letter(c))
        return (d.width if d is not None and d.width else dw)
    box_w = sum(int(cw(c) * 7 + 5) for c in range(c1, c2 + 1))
    box_h = sum(int((ws.row_dimensions[r].height or dh) * 96 / 72) for r in range(r1, r2 + 1))
    img = XLImage(sig_path); ratio = img.width / img.height
    h = min(max_h, box_h - 2 * pad)
    if h * ratio > box_w - 2 * pad:
        h = int((box_w - 2 * pad) / ratio)
    place_signature(ws, sig_path, f"{get_column_letter(c1)}{r1}", max(12, int(h)), x_off_px=pad + 3, y_off_px=max(2, (box_h - h) // 2))


def place_signature(ws, sig_path, cell, height_px, x_off_px=8, y_off_px=4):
    """Put a signature image in a signature box.
    Anchored at both corners so later row-height changes cannot move it out of its box."""
    img = XLImage(sig_path)
    ratio = img.width / img.height
    img.height = height_px; img.width = int(height_px * ratio)
    col0 = column_index_from_string(re.match(r"[A-Z]+", cell).group())
    row0 = int(re.search(r"\d+", cell).group())
    default_w = ws.sheet_format.defaultColWidth or 8.43
    default_h = ws.sheet_format.defaultRowHeight or 15
    def col_px(c):
        d = ws.column_dimensions.get(get_column_letter(c))
        return int(((d.width if d is not None and d.width else default_w) * 7) + 5)
    def row_px(r):
        h = ws.row_dimensions[r].height or default_h
        return int(h * 96 / 72)
    # walk right / down from the top-left cell to find where the picture ends
    c, remaining = col0, x_off_px + img.width
    while remaining > col_px(c):
        remaining -= col_px(c); c += 1
    r, remaining_r = row0, y_off_px + img.height
    while remaining_r > row_px(r):
        remaining_r -= row_px(r); r += 1
    anchor = TwoCellAnchor(editAs="oneCell")
    anchor._from = AnchorMarker(col=col0 - 1, colOff=pixels_to_EMU(x_off_px), row=row0 - 1, rowOff=pixels_to_EMU(y_off_px))
    anchor.to = AnchorMarker(col=c - 1, colOff=pixels_to_EMU(remaining), row=r - 1, rowOff=pixels_to_EMU(remaining_r))
    img.anchor = anchor
    ws.add_image(img)

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, 'blank_master.xlsx')
CONTRACT = 'PMS CONTRACT-OTHERS (OPENED (2'
MONTH_ABBR = ['JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP','OCT','NOV','DEC']

# Performance expectations (objectives) supplied for UCH clinical nursing
PE = [None,
 "Participates in direct holistic patient care and delivers individualized nursing care.",
 "Provides highest possible quality care by; Planning, organizing, staffing, directing, coordinating and evaluating all factors contributing to nursing care.",
 "Provides instant nursing care in emergency situation and informs physician.",
 "Inspects the environment and act to maintain excellent hygiene and safety.",
 "Leads, directs and guide other nursing staff in delivering nursing care.",
 "Compiles and maintains accurate statistics of nursing personnel.",
 "Supervises all nursing activities in the ward / unit.",
 "Applies and implements hospital policies, rules and regulations.",
 "Organizes periodic Ward Orientation sessions for practicing Nurses and Educating them on new trends in clinical nursing practice.",
 "Participates in activities of education, investigation, teaching and research.",
 "Ensures high quality and cost-effective nursing care of all patients.",
 "Liaises with schools through the Quality Improvement unit/Continuing Education in respect of student nurses deployed for clinical experience / exposure.",
 "Taking charge of the Ward / Unit in absence of the Assistant Director/ Deputy Director of Nursing.",
 "Promote nurses' welfare.",
 "Create positive work environment that will promote harmonious relationship.",
 "Foster a supportive and compassionate environment to care for patients and their families.",
 "Prepares annual comprehensive budget proposal for assigned Ward / Unit.",
 "Prepares and writes monthly, quarterly and annual reports on areas of responsibilities.",
 "Expands knowledge and capability through Continuing Education and Professional Development programmes.",
 "Supervises closely and adequately Nurse Interns undergoing internship as well as students and Health Assistants.",
 "Co-ordinates the appraisal of staff within the assigned ward/ Unit.",
 "Inspects food to ensure that the food is appropriate for patient consumption.",
 "Mentors nurses.",
 "Participates in conflict management/ resolution.",
 "Participates in professional association.",
 "Participates in research projects as it affects nursing practice.",
 "The Nurse performs duties harmoniously with other Health Care Practitioners and generally embraces team approach (Team Work).",
 "Accepts other assignment(s) as delegated.",
]
# keyword -> preferred objectives, so each KRA gets a fitting expectation
PE_RULES = [
 (r'emergenc|prompt|respon|triage|critical|acute', [3, 1, 7, 13]),
 (r'hygien|disinfect|steril|clean|infect|wash|asepti|decontam|ipc', [4, 7, 11]),
 (r'train|teach|study|capacity|orient|\bsop\b|standard operating|mentor|student|intern|clinical rotation|osce', [9, 19, 23, 20, 12, 10]),
 (r'document|record|report|statistic|ecm|data|audit|dashboard', [18, 6, 26, 7]),
 (r'outreach|communit|screen|health information|health talk|educat|awareness|enlighten|rall', [10, 16, 25, 27, 28]),
 (r'remind|specialist', [27, 16, 2]),
 (r'workflow|appointment|schedul|theatre|clinic|waiting|referral|wait time|bed management', [2, 5, 13, 11, 17]),
 (r'administ|committee|polic|pms|promotion|apprais|recruit|induction|memo', [8, 21, 6, 28, 17]),
 (r'staff|welfare|roster|collaborat|team', [14, 15, 5, 27, 24]),
 (r'mother|baby|babies|deliver|antenatal|patient care|care deliver|pain|palliat|symptom', [1, 16, 11, 3]),
]
DEFAULT_PE = [1, 11, 27, 28, 2, 7, 5]

MONTH_WIDTHS = {'B': 38, 'C': 17, 'D': 30, 'E': 17, 'F': 42}


def norm(s):
    return ' '.join(str(s or '').split()).strip(' .•').lower()


def upper(s):
    return (s or '').strip().upper()


def ippis_val(s):
    s = (s or '').strip()
    return int(s) if s.isdigit() else s


def phone_val(s):
    s = re.sub(r'\s+', '', (s or ''))
    if s.startswith('+234'):
        s = '0' + s[4:]
    return s


def num_val(s):
    """'85' -> 85, '85%' -> 85, '1' -> 1, '' -> None, other text kept."""
    s = (s or '').strip()
    if not s:
        return None
    t = s.replace('%', '').replace(',', '').strip()
    try:
        f = float(t)
        return int(f) if f.is_integer() else f
    except ValueError:
        return s


def full(p, order='surname_first'):
    parts = [p.get('surname'), p.get('first_name'), p.get('other_name')]
    return ' '.join(x.strip() for x in parts if x and x.strip())


CUSTOM_ROW0 = 1000   # a nurse's own tasks (not on the department list) are staged here, then removed before saving


def _band(b):
    try:
        return float(b) if b is not None and str(b).replace('.', '', 1).isdigit() else b
    except ValueError:
        return b


def apply_own_contract(ws, data):
    """Replace the hospital-wide task list (rows 19-99) with the person's own contract when one is supplied
    (data['own_contract'], e.g. the department's individual contract for clerical staff). Unused rows are hidden."""
    rows = data.get('own_contract') or []
    if not rows or ws.cell(17, 1).value != 'S/N':
        return 0
    for m in [str(m) for m in ws.merged_cells.ranges if 19 <= m.min_row <= 99]:
        ws.unmerge_cells(m)
    for r in range(19, 100):
        for c in range(1, 27):
            ws.cell(r, c).value = None
    r = 19
    groups = []
    for i, row in enumerate(rows[:81]):
        r = 19 + i
        if row.get('code'):
            groups.append([r, r])
        elif groups:
            groups[-1][1] = r
        vals = {1: row.get('code') or None, 2: row.get('kra'), 6: row.get('weight_group') if row.get('code') else None,
                7: row.get('objective'), 11: row.get('weight'), 12: f'=K{r}/$K$100*70', 13: row.get('target'),
                14: row.get('kpi'), 19: row.get('unit')}
        for col, v in vals.items():
            ws.cell(r, col).value = v
        for j, b in enumerate((list(row.get('bands') or []) + [None] * 6)[:6]):
            ws.cell(r, 21 + j).value = _band(b)
        for c in range(1, 27):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            if r > 19:
                cell.font = copy(ws.cell(19, c).font)        # one font per column, as on the first row
        for a, b in ((2, 5), (7, 10), (14, 18), (19, 20)):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    for a, b in groups:
        if b > a:
            for col in (1, 6):
                ws.merge_cells(start_row=a, start_column=col, end_row=b, end_column=col)
    for rr in range(r + 1, 100):
        ws.row_dimensions[rr].hidden = True
    return len(rows)


def add_custom_tasks(pms, data):
    """Write tasks that carry their own target/scale ("custom") into scratch rows of the contract sheet,
    so every lookup, the quarterly pages and the summary treat them like department tasks. Returns how many.
    A person's own contract (data['own_contract']) is put in place first."""
    apply_own_contract(pms, data)
    n = 0
    for k in data.get('kras') or []:
        c = k.get('custom') if isinstance(k, dict) else None
        if not c:
            continue
        r = CUSTOM_ROW0 + n; n += 1
        bands = (c.get('bands') or [None] * 6) + [None] * 6
        vals = {1: k.get('code') or 'OWN', 2: k.get('kra'), 6: c.get('weight'), 7: c.get('objective'), 11: c.get('weight'),
                13: c.get('target'), 14: c.get('kpi'), 19: c.get('unit') or '%'}
        for col, v in vals.items():
            pms.cell(r, col).value = v
        for j in range(6):
            b = bands[j]
            try:
                b = float(b) if b is not None and str(b).replace('.', '', 1).isdigit() else b
            except ValueError:
                pass
            pms.cell(r, 21 + j).value = b
    return n


class Contract:
    """Lookup of KRA rows in the performance contract sheet (department tasks, then a nurse's own)."""
    def __init__(self, ws):
        self.ws = ws
        self.rows = {}
        code = ''
        for r in list(range(19, 100)) + list(range(CUSTOM_ROW0, CUSTOM_ROW0 + 50)):
            if r == CUSTOM_ROW0:
                code = ''
            if r >= CUSTOM_ROW0 and (ws._cells.get((r, 2)) is None or not ws._cells[(r, 2)].value):
                continue                                   # look without creating cells in the scratch area
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a not in (None, ''):
                code = str(a).strip()
                if code.replace('.', '').isdigit():
                    code = str(int(float(code)))
            if b:
                if r >= CUSTOM_ROW0:                           # a nurse's own wording wins over a department row
                    self.rows[(code, norm(b))] = r; self.rows[('*', norm(b))] = r; continue
                self.rows[(code, norm(b))] = r
                self.rows.setdefault(('*', norm(b)), r)

    def find(self, code, text):
        return self.rows.get((code, norm(text))) or self.rows.get(('*', norm(text)))

    def code_for_row(self, r):
        while r > 18 and not self.ws.cell(r, 1).value:
            r -= 1
        v = self.ws.cell(r, 1).value
        s = str(v).strip()
        return str(int(float(s))) if s.replace('.', '').isdigit() else s


def pick_objectives(month_rows):
    used, out = set(), []
    for mrows in month_rows:
        chosen = []
        for row in mrows:
            if not row['kra']:
                chosen.append(None); continue
            text = norm(row['kra'])
            cands = []
            for pat, prefs in PE_RULES:
                if re.search(pat, text):
                    cands += prefs
            cands += DEFAULT_PE + list(range(1, len(PE)))
            pick = next((c for c in cands if c not in used), None)
            if pick is None:           # all 28 used: reuse the best fit
                pick = cands[0]
            used.add(pick); chosen.append(pick)
        out.append(chosen)
    return out


M_MAX = 10       # most tasks reported in one month (the monthly page has rows 12-22)
Q_ROWS = 6       # task rows printed on a blank quarterly appraisal (19-24)
Q_MAX = 10       # most tasks one quarterly appraisal will take; extra rows are inserted above the totals row


def expand_quarter(ws, extra):
    """Make room for `extra` more task rows on a quarterly appraisal: insert them above the totals row (25) and move
    everything below down, keeping merged boxes, row heights and formulas pointing at the right cells."""
    if extra <= 0:
        return
    below = sorted([(r, d.height) for r, d in ws.row_dimensions.items() if r >= 25 and d.height is not None], reverse=True)
    merged = [m for m in ws.merged_cells.ranges if m.min_row >= 25]
    for m in merged:
        ws.merged_cells.remove(m)
    ws.insert_rows(25, extra)
    for m in merged:
        m.shift(0, extra)
        ws.merged_cells.add(m)
    for r, h in below:
        ws.row_dimensions[r + extra].height = h
        ws.row_dimensions[r].height = None
    ref = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")

    def shift(f):
        f = ref.sub(lambda m: m.group(1) + str(int(m.group(2)) + extra) if int(m.group(2)) >= 25 else m.group(0), f)
        return re.sub(r"([A-Z]{1,3})19:([A-Z]{1,3})24\b", lambda m: f"{m.group(1)}19:{m.group(2)}{24 + extra}", f)
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and not isinstance(c, MergedCell):
                c.value = shift(c.value)


def order_sheets(wb):
    """Tabs in reading order: the closed contracts, the open contract, then Jan, Feb, Mar, Q1, Apr, May, Jun, Q2 ... Dec, Q4."""
    def key(i_ws):
        i, ws = i_ws
        n = ws.title.strip().upper()
        if n.startswith('PMS CONTRACT'):
            return (0, 1 if 'OPENED' in n else 0, i)
        for mi, m in enumerate(MONTH_ABBR):
            if n.startswith(m + ' '):
                return (1, mi, 0)
        q = re.match(r'Q([1-4]) ', n)
        if q:
            return (1, int(q.group(1)) * 3 - 1, 1)      # right after the quarter's third month
        return (2, 0, i)
    active = wb.active.title if wb.active is not None else None
    wb._sheets = [ws for _, ws in sorted(enumerate(wb._sheets), key=key)]
    if active:
        wb.active = wb._sheets.index(wb[active])
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.active


def fill(data, out_path, signature=None, supervisor_signature=None, cso_signature=None):
    year = int(data.get('year') or 2025)
    emp, sup, cso = data.get('employee', {}), data.get('supervisor', {}), data.get('countersigning_officer', {})
    extras = data.get('extras', {}) or {}
    wb = openpyxl.load_workbook(TEMPLATE)
    pms = wb[CONTRACT]
    n_custom = add_custom_tasks(pms, data)
    if 'PMS CONTRACT-OTHERS (CLOSED)' in wb.sheetnames:
        apply_own_contract(wb['PMS CONTRACT-OTHERS (CLOSED)'], data)
    contract = Contract(pms)
    notes = []

    # ---- contract page: people
    vals = {
        'C7': emp.get('surname'), 'H7': emp.get('first_name'), 'M7': emp.get('other_name'), 'R7': emp.get('designation'),
        'C8': ippis_val(emp.get('ippis')), 'H8': emp.get('email'), 'M8': phone_val(emp.get('phone')), 'R8': emp.get('department') or 'Clinical Nursing',
        'C10': sup.get('surname'), 'H10': sup.get('first_name'), 'M10': sup.get('other_name'), 'R10': sup.get('designation'),
        'C11': ippis_val(sup.get('ippis')), 'H11': sup.get('email'), 'M11': phone_val(sup.get('phone')),
        'C13': cso.get('surname'), 'H13': cso.get('first_name'), 'M13': cso.get('other_name'), 'R13': cso.get('designation'),
        'C14': ippis_val(cso.get('ippis')), 'H14': cso.get('email'), 'M14': phone_val(cso.get('phone')),
    }
    for a, v in vals.items():
        pms[a] = v if v not in ('',) else None
    pms['C5'] = datetime(year, 1, 1); pms['G5'] = datetime(year, 12, 31)

    # ---- monthly reviews
    monthly = (data.get('monthly') or [])[:12]
    while len(monthly) < 12:
        monthly.append({'rows': []})
    month_rows = []
    for m in monthly:
        rows = [r for i, r in enumerate((m.get('rows') or [])[:M_MAX]) if i < 2 or (r or {}).get('kra')]   # tasks 3+ only if used
        while len(rows) < 2:
            rows.append({})
        month_rows.append([{'code': r.get('code', ''), 'kra': (r.get('kra') or '').strip(), 'objective': (r.get('objective') or '').strip(),
                            'output': r.get('output', ''), 'issues': (r.get('issues') or '').strip()} for r in rows])
    objectives = pick_objectives(month_rows)
    sys.path.insert(0, HERE)
    import review
    month_entries = review.monthly_entries(month_rows, contract, pms)
    review.contract_comments(pms)
    month_sheets = {}
    for name in wb.sheetnames:
        if 'MONTHLY' in name:
            month_sheets[MONTH_ABBR.index(name[:3])] = name
    sup_name = upper(' '.join(x for x in [sup.get('surname'), sup.get('first_name'), sup.get('other_name')] if x))
    for mi in range(12):
        ws = wb[month_sheets[mi]]
        for col, w in MONTH_WIDTHS.items():
            ws.column_dimensions[col].width = w
        start = datetime(year, mi + 1, 1); end = datetime(year, mi + 1, calendar.monthrange(year, mi + 1)[1])
        hdr = {'B5': upper(emp.get('surname')), 'E5': upper(emp.get('other_name')), 'B6': upper(emp.get('first_name')),
               'E6': upper(emp.get('department') or 'Clinical Nursing'), 'B7': ippis_val(emp.get('ippis')),
               'E7': sup_name, 'B8': upper(emp.get('designation')), 'E8': upper(sup.get('designation')), 'B9': start, 'E9': end}
        for a, v in hdr.items():
            ws[a] = v or None
        ws['B9'].number_format = ws['E9'].number_format = 'm/d/yyyy'
        filled_any = False
        for ri, row in enumerate(month_rows[mi]):
            r = 12 + ri
            if not row['kra'] and num_val(row['output']) is None:
                continue
            filled_any = True
            ws.cell(r, 1).value = row['kra'] or None
            crow_o = contract.find(row['code'], row['kra']) if row['kra'] else None
            own_obj = (' '.join(str(pms.cell(crow_o, 7).value or '').split()) or None) \
                if crow_o and (crow_o >= CUSTOM_ROW0 or data.get('objectives_from_contract')) else None
            ws.cell(r, 2).value = row['objective'] or own_obj or (PE[objectives[mi][ri]] if objectives[mi][ri] else None)
            ws.cell(r, 3).value = start; ws.cell(r, 4).value = end
            ws.cell(r, 3).number_format = ws.cell(r, 4).number_format = 'm/d/yyyy'
            out = num_val(row['output'])
            ws.cell(r, 5).value = out
            crow = contract.find(row['code'], row['kra']) if row['kra'] else None
            target = pms.cell(crow, 13).value if crow else None
            if isinstance(out, (int, float)) and isinstance(target, (int, float)) and target <= 1 and out <= 1:
                ws.cell(r, 5).number_format = '0%'   # e.g. 1 = 100%
            ws.cell(r, 6).value = row['issues'] or None
            for c in range(1, 7):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        if not filled_any:
            notes.append(f"{MONTH_ABBR[mi].title()}: no figures submitted")
        es = [e for e in month_entries[mi] if e.value is not None]
        if es:
            appraisee_c, appraiser_c = review.monthly_comments(es, [r['issues'] for r in month_rows[mi]], mi, emp.get('first_name'))
            own = ((monthly[mi] or {}).get('appraisee_comment') or '').strip()
            own_sup = ((monthly[mi] or {}).get('appraiser_comment') or '').strip()
            ws['B39'] = own or appraisee_c; ws['B40'] = own_sup or appraiser_c
            for a in ('B39', 'B40'):
                ws[a].alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        for a, key in [('A24', 'outstanding_performance'), ('A27', 'areas_of_improvement'), ('A30', 'training_needs'),
                       ('A34', 'future_goals'), ('A37', 'other_feedback')]:
            v = (extras.get(key) or '').strip()
            ws[a] = v or None
            if v:
                ws[a].alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')

    # ---- quarterly appraisals
    q_sheets = sorted([n for n in wb.sheetnames if 'APPRAISAL' in n], key=lambda n: n.strip()[:2])
    qinfo = {
        'C7': emp.get('surname'), 'E7': emp.get('first_name'), 'G7': emp.get('other_name'), 'I7': emp.get('designation'),
        'C8': ippis_val(emp.get('ippis')), 'E8': emp.get('email'), 'G8': phone_val(emp.get('phone')), 'I8': emp.get('department') or 'Clinical Nursing',
        'C10': sup.get('surname'), 'E10': sup.get('first_name'), 'G10': sup.get('other_name'), 'I10': sup.get('designation'),
        'C11': ippis_val(sup.get('ippis')), 'E11': sup.get('email'), 'G11': phone_val(sup.get('phone')),
        'C13': cso.get('surname'), 'E13': cso.get('first_name'), 'G13': cso.get('other_name'), 'I13': cso.get('designation'),
        'C14': ippis_val(cso.get('ippis')), 'E14': cso.get('email'), 'G14': phone_val(cso.get('phone')),
    }
    periods = [('01/01', '31/03'), ('01/04', '30/06'), ('01/07', '30/09'), ('01/10', '31/12')]
    reset = {month_sheets[mi]: tuple(range(12, 12 + max(6, len(month_rows[mi])))) + (24, 27, 30, 34, 37, 39, 40) for mi in range(12)}
    qoff = {}
    for qi, name in enumerate(q_sheets):
        ws = wb[name]
        ws['D5'] = f"{periods[qi][0]}/{year} TO {periods[qi][1]}/{year}"
        ws.column_dimensions['A'].width = 19; ws.column_dimensions['B'].width = 36
        ws.column_dimensions['F'].width = 19   # roomier supervisor signature box (F59:F61)
        for a, v in qinfo.items():
            ws[a] = v or None
            ws[a].alignment = Alignment(wrap_text=True, vertical='center', horizontal=ws[a].alignment.horizontal)
        seen, ks = set(), []
        for mi in range(qi * 3, qi * 3 + 3):
            for row in month_rows[mi]:
                key = (row['code'], norm(row['kra']))
                if row['kra'] and key not in seen:
                    seen.add(key); ks.append(row)
        if len(ks) > Q_MAX:
            notes.append(f"Q{qi+1}: {len(ks)} different tasks, only the first {Q_MAX} fit on the appraisal page")
            ks = ks[:Q_MAX]
        off = max(0, len(ks) - Q_ROWS)
        expand_quarter(ws, off)
        qoff[name] = off
        for i, row in enumerate(ks):
            r = 19 + i
            if r > 19:
                for c in range(1, 19):
                    s, d = ws.cell(19, c), ws.cell(r, c)
                    d.font, d.border, d.fill, d.number_format = copy(s.font), copy(s.border), copy(s.fill), s.number_format
            cr = contract.find(row['code'], row['kra'])
            ws.cell(r, 2).value = row['kra']
            if cr:
                ws.cell(r, 1).value = contract.code_for_row(cr)
                ws.cell(r, 3).value = pms.cell(cr, 11).value
                ws.cell(r, 3).number_format = ws.cell(r, 5).number_format
                ws.cell(r, 4).value = ' '.join(str(pms.cell(cr, 7).value or '').split()) or None
                ws.cell(r, 5).value = pms.cell(cr, 11).value
                ws.cell(r, 6).value = f'=E{r}/$E${25 + off}*70'
                ws.cell(r, 7).value = ' '.join(str(pms.cell(cr, 14).value or '').split()) or None
                ws.cell(r, 8).value = pms.cell(cr, 13).value
                if isinstance(ws.cell(r, 8).value, (int, float)) and ws.cell(r, 8).value <= 1 and str(pms.cell(cr, 19).value).strip() == '%':
                    ws.cell(r, 8).number_format = '0%'
                ws.cell(r, 9).value = pms.cell(cr, 19).value
                for j in range(6):
                    ws.cell(r, 10 + j).value = pms.cell(cr, 21 + j).value
            else:
                ws.cell(r, 1).value = 'OWN'
                notes.append(f"Q{qi+1}: task not in the contract, weights and KPIs left blank: {row['kra'][:60]}")
            for c in range(1, 16):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='center',
                                                    horizontal='center' if c in (1, 3, 5, 6, 8, 9) or c >= 10 else 'left')
        ws.cell(25 + off, 3).value = f'=SUM(C19:C{24 + off})'
        by_key = {}
        for mi in range(qi * 3, qi * 3 + 3):
            for e in month_entries[mi]:
                by_key.setdefault((e.row['code'], norm(e.row['kra'])), []).append(e)
        review.fill_quarter(ws, qi, ks, by_key, contract, pms, data, year, off)
        reset[name] = (5, 7, 8, 10, 11, 13, 14) + tuple(range(19, 25 + off)) + tuple(r + off for r in list(range(30, 45)) + [52, 54, 55, 56, 57, 58])

    # signatures (only those supplied and approved): appraisee, supervisor/appraiser, counter-signing officer
    if signature or supervisor_signature or cso_signature:
        # give every row an explicit height so no program re-computes positions differently
        for ws in wb.worksheets:
            dh = ws.sheet_format.defaultRowHeight or 15
            for r in range(1, ws.max_row + 1):
                if ws.row_dimensions[r].height is None:
                    ws.row_dimensions[r].height = dh
        # taller signature rows so a signature prints at a readable size (about 2 cm high)
        SIG_PT = 66
        for mi in range(12):
            wb[month_sheets[mi]].row_dimensions[42].height = SIG_PT
        for r in (124, 125, 126):
            pms.row_dimensions[r].height = SIG_PT / 3
        for name in q_sheets:
            for r in (59, 60, 61):
                wb[name].row_dimensions[r + qoff[name]].height = SIG_PT / 3
        big = int(SIG_PT * 96 / 72) - 10                     # pixels: box height less a small margin
        if signature:
            place_signature_fit(pms, signature, 'C124:G126', max_h=big)
            for mi in range(12):
                place_signature_fit(wb[month_sheets[mi]], signature, 'B42:C42', max_h=big)
            for name in q_sheets:
                place_signature_fit(wb[name], signature, f'C{59 + qoff[name]}:D{61 + qoff[name]}', max_h=big)
        if supervisor_signature:
            place_signature_fit(pms, supervisor_signature, 'I124:M126', max_h=big)
            for mi in range(12):
                place_signature_fit(wb[month_sheets[mi]], supervisor_signature, 'E42:F42', max_h=big)
            for name in q_sheets:
                place_signature_fit(wb[name], supervisor_signature, f'F{59 + qoff[name]}:F{61 + qoff[name]}', max_h=big)
        if cso_signature:
            place_signature_fit(pms, cso_signature, 'P124:T126', max_h=big)
            for name in q_sheets:
                place_signature_fit(wb[name], cso_signature, f'H{59 + qoff[name]}:J{61 + qoff[name]}', max_h=big)

    if n_custom:
        pms.delete_rows(CUSTOM_ROW0, 60)                   # scratch rows are not part of the workbook
        for r in [r for r in pms.row_dimensions if r >= CUSTOM_ROW0 - 5]:
            del pms.row_dimensions[r]
    cc = data.get('contract') or {}
    for a, k in (('C120', 'appraisee_comment'), ('I120', 'supervisor_comment'), ('P120', 'cso_comment')):
        v = (cc.get(k) or '').strip()
        if v:
            pms[a] = v
    order_sheets(wb)
    stage = out_path + '.stage.xlsx'
    wb.save(stage)
    sys.path.insert(0, HERE)
    from autofit import autofit
    autofit(stage, out_path, reset)
    os.remove(stage)
    return notes


def signature_from_data(data, path, field='signature'):
    """Write an approved signature (a PNG data URL saved by the form) to `path`; None if there isn't one.
    field: 'signature' (nurse), 'supervisor_signature' or 'cso_signature'."""
    import base64
    sig = data.get(field) or {}
    png = sig.get('png') if isinstance(sig, dict) and sig.get('approved') else None
    prefix = 'data:image/png;base64,'
    if not (isinstance(png, str) and png.startswith(prefix)):
        return None
    raw = base64.b64decode(png[len(prefix):])
    if not raw.startswith(b'\x89PNG'):
        return None
    with open(path, 'wb') as fh:
        fh.write(raw)
    return path


def out_name(data):
    e = data.get('employee', {})
    n = '_'.join(x for x in [e.get('surname'), e.get('first_name')] if x) or 'Nurse'
    n = re.sub(r'[^A-Za-z0-9_-]+', '-', n)
    return f"EPMS_{data.get('year', 2025)}_{n}.xlsx"


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='+')
    ap.add_argument('--out', default='out')
    ap.add_argument('--signature', default=None, help='PNG of the appraisee\'s own signature (transparent background)')
    ap.add_argument('--supervisor-signature', default=None, help='PNG of the supervisor\'s signature')
    ap.add_argument('--cso-signature', default=None, help='PNG of the counter-signing officer\'s signature')
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    for f in a.files:
        d = json.load(open(f))
        d = d.get('data', d)
        p = os.path.join(a.out, out_name(d))
        given = {'signature': a.signature, 'supervisor_signature': a.supervisor_signature, 'cso_signature': a.cso_signature}
        sigs, temp = {}, []
        for field, path in given.items():
            if not path:
                path = signature_from_data(d, f"{p}.{field}.png", field)
                if path:
                    temp.append(path)
            sigs[field] = path
        notes = fill(d, p, sigs['signature'], sigs['supervisor_signature'], sigs['cso_signature'])
        for t in temp:
            os.remove(t)
        print(json.dumps({'file': p, 'signed': {k: bool(v) for k, v in sigs.items()}, 'notes': notes}))
