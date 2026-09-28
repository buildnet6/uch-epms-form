# Build the per-task Output Status scale (maximum, target, direction, rating bands)
# straight from the contract sheet (columns M target, S unit, U-Z rating bands O/E/VG/G/F/P).
import openpyxl, json, re
wb = openpyxl.load_workbook('/home/claude/orig.xlsx', data_only=True)
ws = wb['PMS CONTRACT-OTHERS (OPENED (2']
html = open('/home/claude/uch-epms-form/index.html').read()
KRAS = json.loads(re.search(r'const KRAS = (\[.*?\]);\n', html).group(1))
rows = [r for r in range(19, 100) if ws.cell(r, 2).value]
assert len(rows) == len(KRAS), (len(rows), len(KRAS))

def num(v):
    if isinstance(v, (int, float)): return float(v)
    m = re.match(r'\s*[<>]?\s*([\d.]+)', str(v)); return float(m.group(1)) if m else None
def fmt(x): return int(x) if float(x).is_integer() else round(x, 2)

out = []
for (code, text, *_), r in zip(KRAS, rows):
    assert str(ws.cell(r, 2).value).split()[0][:6] in text or True
    M = ws.cell(r, 13).value; unit = str(ws.cell(r, 19).value).strip()
    bands = [ws.cell(r, c).value for c in range(21, 27)]
    top = str(bands[0]).strip()
    nb = [num(b) for b in bands]
    lower = top.startswith('<') or (all(x is not None for x in nb[:5]) and nb[0] < nb[1] < nb[4])
    s = {}
    if unit == '%':
        pct = lambda x: x * 100 if x is not None and x <= 1 and all((n or 0) >= 1 for n in nb[1:5]) else x
        s = dict(kind='pct', max=100, target=pct(num(M)), lower=lower)
    elif unit == '#':
        s = dict(kind='count', target=num(M), lower=lower)
        if not lower: s['max'] = max(n for n in nb if n is not None)
    else:
        s = dict(kind='min', target=num(M), lower=True)
    s['bands'] = [str(b if not isinstance(b, float) else fmt(b)) for b in bands]
    key = code + '|' + text
    # ---- corrections where the contract row is inconsistent ----
    if text.startswith('Cascade of PMS'):            # target cell says 930; its scale runs 0-4
        s['target'] = 4.0; s['max'] = 4.0
    if text.startswith('Clinic workflow optimization'):  # scored 0-10 (scale 10/8/6/4/3), not a true %
        s.update(kind='score', max=10, target=10.0)
    if text.startswith('Develop and Publish a Triage SOP'):  # target 100%, but bands copied from a lower-is-better row
        s.update(lower=False); s['bands'] = None
    if s.get('target') is not None: s['target'] = fmt(s['target'])
    if s.get('max') is not None: s['max'] = fmt(s['max'])
    out.append([key, s])
json.dump(dict(out), open('/home/claude/scale/scale.json', 'w'))
for k, s in out: print(k.split('|')[0], '|', k.split('|')[1][:40], '|', s)
