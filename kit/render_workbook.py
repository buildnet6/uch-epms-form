#!/usr/bin/env python3
"""Render a filled EPMS workbook the way it will look when printed/submitted.

Makes a print-ready copy (only the nurse's own sheets, each fitted to the page width, landscape where the form is wide),
converts it to PDF with LibreOffice and, optionally, to one JPEG per page for on-screen viewing.
The original workbook is never changed.

    python3 render_workbook.py EPMS_2025_Name.xlsx --out outdir [--images]
"""
import os, re, sys, json, shutil, argparse, subprocess, tempfile
import openpyxl
from openpyxl.worksheet.properties import PageSetupProperties

CONTRACT = 'PMS CONTRACT-OTHERS (OPENED (2'
MONTHS = ['JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC']


def order_key(name):
    n = name.strip().upper()
    if n == CONTRACT.upper():
        return (0, 0)
    for i, m in enumerate(MONTHS):
        if n.startswith(m + ' '):
            return (1, i)
    q = re.match(r'Q([1-4]) ', n)
    if q:
        return (1, int(q.group(1)) * 3 - 0.5)       # each quarter right after its third month
    return (9, 0)


def label(name):
    n = name.strip().upper()
    if n == CONTRACT.upper():
        return 'Performance contract'
    for i, m in enumerate(MONTHS):
        if n.startswith(m + ' '):
            return ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September',
                    'October', 'November', 'December'][i] + ' review'
    q = re.match(r'Q([1-4]) ', n)
    if q:
        return f'Quarter {q.group(1)} appraisal'
    return name.strip()


def last_used_row(ws, max_col=26):
    last = 1
    for (r, c), cell in ws._cells.items():
        if c <= max_col and cell.value not in (None, ''):
            last = max(last, r)
    for img in getattr(ws, '_images', []):
        try:
            last = max(last, img.anchor._from.row + 4)
        except Exception:
            pass
    return last


def sheet_size(ws, max_col, last_row):
    """Approximate printed width/height in points (column width units ~7pt, default row 15pt)."""
    w = sum((ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.43) * 7 for c in range(1, max_col + 1)
            if not ws.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden)
    h = sum((ws.row_dimensions[r].height or 15) for r in range(1, last_row + 1) if not ws.row_dimensions[r].hidden)
    return w, h


def setup_page(ws, whole=True):
    """A4, fitted to the page width; a whole form on one page (monthly/quarterly), the long contract over several."""
    max_col = ws.max_column if ws.max_column < 40 else 26
    last = last_used_row(ws, max_col)
    w, h = sheet_size(ws, max_col, last)
    ps = ws.page_setup
    ps.paperSize = ws.PAPERSIZE_A4
    ps.orientation = 'portrait' if whole else 'landscape'     # single forms are tall; the contract is wide
    ps.fitToWidth, ps.fitToHeight = 1, (1 if whole else 0)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.print_area = f"A1:{openpyxl.utils.get_column_letter(max_col)}{last}"
    ws.print_title_rows = None
    return ps.orientation


def print_copy(src, dst):
    wb = openpyxl.load_workbook(src)
    keep = [ws for ws in wb.worksheets if order_key(ws.title)[0] < 9]
    for ws in wb.worksheets:
        if ws not in keep:
            ws.sheet_state = 'hidden'
    for ws in keep:
        setup_page(ws, whole=order_key(ws.title)[0] == 1)
    # print in reading order: contract, then each month with its quarter after the third month
    wb._sheets = sorted(keep, key=lambda w: order_key(w.title)) + [w for w in wb.worksheets if w not in keep]
    keep = wb._sheets[:len(keep)]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    keep[0].sheet_view.tabSelected = True
    wb.save(dst)
    return [(ws.title, label(ws.title)) for ws in sorted(keep, key=lambda w: order_key(w.title))]


def render(src, outdir, images=False, dpi=120):
    os.makedirs(outdir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src))[0]
    with tempfile.TemporaryDirectory() as tmp:
        cp = os.path.join(tmp, base + '.xlsx')
        sheets = print_copy(src, cp)
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'pdf', '--outdir', tmp, cp],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        pdf = os.path.join(outdir, base + '_Rendered.pdf')
        shutil.move(os.path.join(tmp, base + '.pdf'), pdf)
    pages = []
    if images:
        from pypdf import PdfReader
        n = len(PdfReader(pdf).pages)
        prefix = os.path.join(outdir, base + '_p')
        subprocess.run(['pdftoppm', '-r', str(dpi), '-jpeg', '-jpegopt', 'quality=72', pdf, prefix], check=True)
        pages = sorted(f for f in os.listdir(outdir) if f.startswith(base + '_p') and f.endswith('.jpg'))
        assert len(pages) == n, (len(pages), n)
    # which pages belong to which form: the contract runs over several pages, every other form is one page
    from pypdf import PdfReader
    texts = [(pg.extract_text() or '').lstrip()[:40].upper() for pg in PdfReader(pdf).pages]
    first_form = next((i for i, t in enumerate(texts) if t.startswith('MONTHLY') or t.startswith('PERFORMANCE APPRAISAL')), len(texts))
    spans, at = [], 0
    for title, lab in sheets:
        n = first_form if title == CONTRACT else 1
        spans.append({'sheet': title, 'label': lab, 'first': at + 1, 'pages': n}); at += n
    ok = at == len(texts)
    return {'pdf': pdf, 'sheets': spans, 'pages': pages, 'page_count': len(texts), 'map_ok': ok}


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='+')
    ap.add_argument('--out', default='out')
    ap.add_argument('--images', action='store_true')
    a = ap.parse_args()
    for f in a.files:
        print(json.dumps(render(f, a.out, a.images)))
